from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs' / '使用手册.md'
target = root / 'docs' / '使用手册.docx'

lines = source.read_text(encoding='utf-8').splitlines()

doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
normal.font.size = Pt(10.5)

def set_style_heading(paragraph):
    for run in paragraph.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_table(rows):
    if not rows:
        return
    rows = [row for row in rows if row]
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = 'Light Grid Accent 1'
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(row):
            cells[col].text = value if col < len(row) else ''
            if index == 0:
                for paragraph in cells[col].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

table_rows = []
i = 0
while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    if stripped.startswith('|'):
        cells = [cell.strip() for cell in stripped.strip('|').split('|')]
        if not all(set(cell) <= set('-: ') for cell in cells if cell):
            table_rows.append(cells)
        i += 1
        continue
    else:
        add_table(table_rows)
        table_rows = []

    if stripped.startswith('# '):
        doc.add_heading(stripped[2:].strip(), level=1)
    elif stripped.startswith('## '):
        doc.add_heading(stripped[3:].strip(), level=2)
    elif stripped.startswith('### '):
        doc.add_heading(stripped[4:].strip(), level=3)
    elif stripped.startswith('> '):
        paragraph = doc.add_paragraph(stripped[2:].strip())
        paragraph.style = doc.styles['Quote']
    elif stripped.startswith('- '):
        doc.add_paragraph(stripped[2:].strip(), style='List Bullet')
    elif len(stripped) > 2 and stripped[0].isdigit() and stripped[1] == '.':
        doc.add_paragraph(stripped.split('.', 1)[1].strip(), style='List Number')
    elif stripped:
        doc.add_paragraph(stripped)
    i += 1

if table_rows:
    add_table(table_rows)

doc.save(target)
print(target)
