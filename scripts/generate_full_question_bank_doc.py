from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from question_bank import ROLES


OUTPUT = Path.home() / "Desktop" / "AI+XR业务岗位六维能力行为评价题库（七岗位210题正式版）.docx"

BLUE = "1178C5"
BLUE_DARK = "06345A"
BLUE_MID = "D9ECF8"
BLUE_LIGHT = "EDF7FD"
BLUE_PALE = "F7FBFE"
TEXT = "294B63"
GREY = "6C8395"
LINE = "B7D8EB"
WHITE = "FFFFFF"

CATEGORIES = ("通用基础能力", "专业核心能力", "职业素养能力")
EXPECTED_DIMENSIONS = {
    "销售": ("市场判断", "沟通执行", "拓展客户", "推动成交", "诚信服务", "抗压担当"),
    "商务": ("逻辑表达", "沟通协调", "挖掘需求", "方案签约", "严谨风控", "客户意识"),
    "项目经理": ("计划统筹", "协调推进", "进度质量", "资源成本", "结果负责", "抗压复盘"),
    "技术交付": ("理解执行", "响应改进", "技术制作", "质量交付", "专注负责", "协作成长"),
    "线下运营": ("现场统筹", "沟通应变", "开店带队", "运营增收", "服务意识", "抗压改进"),
    "IP运营": ("创意策划", "统筹分析", "IP运营", "产品变现", "审美创新", "客户意识"),
    "内容发行": ("市场判断", "资源沟通", "内容适配", "渠道落地", "目标执行", "长期经营"),
}


def style_run(run, *, size=None, bold=None, color=None, font="微软雅黑") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, **style):
    run = paragraph.add_run(text)
    style_run(run, **style)
    return run


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, "—  ", size=8.5, color=GREY)
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_begin, instr, fld_separate, fld_end))
    style_run(run, size=8.5, color=GREY)
    add_text(paragraph, "  —", size=8.5, color=GREY)


def configure_section(section, header_text: str) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.72)
    section.footer_distance = Cm(0.72)

    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, header_text, size=8.3, color=GREY)

    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    add_page_number(fp)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, size=17 if level == 1 else 13, bold=True, color=BLUE_DARK)
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), BLUE)
        p_bdr.append(bottom)
        p_pr.append(p_bdr)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, size=10.2) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, size=size, bold=True, color=BLUE_DARK)
        add_text(p, text[len(bold_prefix):], size=size, color=TEXT)
    else:
        add_text(p, text, size=size, color=TEXT)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    add_text(p, text, size=9.8, color=TEXT)


def add_info_box(doc: Document, title: str, lines: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, BLUE_PALE)
    set_cell_margins(cell, 150, 180, 150, 180)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 7, "color": LINE},
        bottom={"val": "single", "sz": 7, "color": LINE},
        left={"val": "single", "sz": 18, "color": BLUE},
        right={"val": "single", "sz": 7, "color": LINE},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    add_text(p, title, size=10.5, bold=True, color=BLUE_DARK)
    for label, value in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        add_text(p, f"{label}：", size=9.3, bold=True, color=BLUE_DARK)
        add_text(p, value, size=9.3, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_scale_table(doc: Document) -> None:
    headers = ("选项", "内部百分制", "雷达图五分值", "统一行为含义")
    rows = (
        ("A", "20", "1分", "明显不能胜任；需要持续督促，或由他人代为完成。"),
        ("B", "40", "2分", "只能完成一部分；依赖提醒、指导或他人补救。"),
        ("C", "60", "3分", "基本达到岗位要求；能独立完成约80%的常规工作。"),
        ("D", "80", "4分", "能够稳定、完整地达到岗位要求。"),
        ("E", "100", "5分", "在完整达标基础上，能够主动优化、创造额外价值或形成可复制方法。"),
        ("无法判断", "不计分", "不计分", "没有接触过该行为，或缺少足够的事实依据。"),
    )
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = (Cm(2.1), Cm(2.7), Cm(2.8), Cm(10.1))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=9.0, bold=True, color=WHITE)
    repeat_table_header(table.rows[0])
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for ci, value in enumerate(values):
            cells[ci].width = widths[ci]
            cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[ci], 80, 100, 80, 100)
            if ri % 2 == 0:
                set_cell_shading(cells[ci], BLUE_PALE)
            p = cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci < 3 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=8.8, bold=ci == 0, color=BLUE_DARK if ci < 3 else TEXT)


def add_overview_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("职位编号", "职位", "六个能力维度", "题量")
    widths = (Cm(2.2), Cm(2.7), Cm(11.0), Cm(1.6))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=9, bold=True, color=WHITE)
    repeat_table_header(table.rows[0])
    for ri, role in enumerate(ROLES, 1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        values = (f"{ri:02d} / {role.code}", role.name, " ｜ ".join(d.name for d in role.dimensions), "30题")
        for ci, value in enumerate(values):
            cells[ci].width = widths[ci]
            set_cell_margins(cells[ci], 95, 100, 95, 100)
            if ri % 2 == 1:
                set_cell_shading(cells[ci], BLUE_PALE)
            p = cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci != 2 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=8.7, bold=ci in (0, 1), color=BLUE_DARK if ci in (0, 1) else TEXT)


def add_role_fields(doc: Document, role_name: str) -> None:
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    values = (
        ("被评估人姓名", "________________", "职位", role_name),
        ("评估日期", "____年__月__日", "评分人/层级", "____________ / □上级 □同级 □下属 □其他"),
    )
    for ri, row_values in enumerate(values):
        for ci, value in enumerate(row_values):
            cell = table.rows[ri].cells[ci]
            set_cell_margins(cell, 85, 100, 85, 100)
            if ci % 2 == 0:
                set_cell_shading(cell, BLUE_MID)
            p = cell.paragraphs[0]
            add_text(p, value, size=8.8, bold=ci % 2 == 0, color=BLUE_DARK if ci % 2 == 0 else TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_dimension_banner(doc: Document, dimension_index: int, dimension) -> None:
    first = (dimension_index - 1) * 5 + 1
    last = first + 4
    add_heading(
        doc,
        f"{dimension_index}. {dimension.category}｜{dimension.name}（第{first:02d}—{last:02d}题）",
        2,
    )
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, BLUE_LIGHT)
    set_cell_margins(cell, 100, 130, 100, 130)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 6, "color": LINE},
        bottom={"val": "single", "sz": 6, "color": LINE},
        left={"val": "single", "sz": 6, "color": LINE},
        right={"val": "single", "sz": 6, "color": LINE},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.2
    add_text(p, "维度定义：", size=9.2, bold=True, color=BLUE_DARK)
    add_text(p, dimension.definition, size=9.2, color=TEXT)


def add_question_block(doc: Document, role, dimension_index: int, question_index: int, question) -> None:
    role_question_no = (dimension_index - 1) * 5 + question_index
    code = f"{role.code}-{dimension_index:02d}-{question_index:02d}"
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    row = table.rows[0]
    prevent_row_split(row)
    cell = row.cells[0]
    set_cell_margins(cell, 120, 150, 105, 150)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 14, "color": BLUE},
        bottom={"val": "single", "sz": 6, "color": LINE},
        left={"val": "single", "sz": 6, "color": LINE},
        right={"val": "single", "sz": 6, "color": LINE},
    )

    header = cell.paragraphs[0]
    header.paragraph_format.space_after = Pt(3)
    set_paragraph_shading(header, BLUE_LIGHT)
    add_text(header, f"第{role_question_no:02d}题　{question.focus}", size=10.0, bold=True, color=BLUE_DARK)
    add_text(header, f"　｜　题目编码：{code}", size=8.7, color=GREY)

    prompt = cell.add_paragraph()
    prompt.paragraph_format.space_before = Pt(2)
    prompt.paragraph_format.space_after = Pt(4)
    prompt.paragraph_format.line_spacing = 1.22
    add_text(prompt, question.prompt, size=9.3, bold=True, color=TEXT)

    for label, option in zip(("A", "B", "C", "D", "E"), question.options):
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.28)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.12
        add_text(p, f"□ {label}　", size=8.7, bold=True, color=BLUE)
        add_text(p, option, size=8.7, color=TEXT)

    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.28)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "□ 无法判断　", size=8.6, bold=True, color=GREY)
    add_text(p, "没有足够接触或事实依据（本题不计分）。", size=8.6, color=GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def validate_bank() -> None:
    assert len(ROLES) == 7
    assert set(role.name for role in ROLES) == set(EXPECTED_DIMENSIONS)
    codes: list[str] = []
    prompts: list[str] = []
    for role in ROLES:
        assert len(role.dimensions) == 6
        assert tuple(d.name for d in role.dimensions) == EXPECTED_DIMENSIONS[role.name]
        assert tuple(d.category for d in role.dimensions) == (
            "通用基础能力",
            "通用基础能力",
            "专业核心能力",
            "专业核心能力",
            "职业素养能力",
            "职业素养能力",
        )
        for di, dimension in enumerate(role.dimensions, 1):
            assert dimension.category in CATEGORIES
            assert len(dimension.questions) == 5
            for qi, question in enumerate(dimension.questions, 1):
                code = f"{role.code}-{di:02d}-{qi:02d}"
                codes.append(code)
                prompts.append(question.prompt)
                assert question.focus.strip() and question.prompt.endswith("？")
                assert len(question.options) == 5
                assert all(option.strip() for option in question.options)
                assert question.options[0] != question.options[-1]
    assert len(codes) == 210 and len(codes) == len(set(codes))
    assert len(prompts) == 210 and len(prompts) == len(set(prompts))


def build_document() -> Document:
    validate_bank()
    doc = Document()
    configure_section(doc.sections[0], "AI+XR业务岗位六维能力行为评价题库｜正式版")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10)
    for level in (1, 2):
        style = styles[f"Heading {level}"]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string(BLUE_DARK)

    props = doc.core_properties
    props.title = "AI+XR业务岗位六维能力行为评价题库（七岗位210题正式版）"
    props.subject = "七岗位六维能力行为选择题"
    props.author = "AI+XR业务能力模型项目组"
    props.keywords = "能力评价, 六维能力图, 行为选择题, AI+XR"
    props.comments = "依据已确认的岗位六维能力模型编制。"

    # 封面
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "AI + XR", size=14, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    add_text(p, "业务岗位六维能力", size=26, bold=True, color=BLUE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "行为评价题库", size=26, bold=True, color=BLUE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    add_text(p, "七岗位全量正式版 · 42个维度 · 210道题", size=12.5, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    add_text(p, "版本 V1.0", size=10, bold=True, color=BLUE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "2026年8月", size=9.5, color=GREY)
    doc.add_page_break()

    # 使用说明
    add_heading(doc, "一、文档使用说明", 1)
    add_body(doc, "本题库用于生成不同岗位的六维能力图。每个职位包含6个能力维度，每个维度5题，共30题；七个职位合计210题。")
    add_info_box(
        doc,
        "评价原则",
        [
            ("评价对象", "评价被评估人在指定周期内实际表现出来的行为，不评价性格、意愿或单次偶然结果"),
            ("证据要求", "优先依据具体项目、客户、数据、交付物或现场事实作答"),
            ("选择方法", "选择最接近其通常表现的一个选项；若没有足够接触，请选择“无法判断”"),
            ("评价周期", "以各职位章节标注的周期为准"),
        ],
    )
    add_heading(doc, "二、统一评分标尺", 1)
    add_body(doc, "所有题目使用同一套五档行为标尺。选项文字已经按具体工作场景展开，评分人只需选择最符合事实的一项。")
    add_scale_table(doc)
    add_body(doc, "说明：题库中的A—E不直接展示分数；系统内部可按20、40、60、80、100计分，并分别转换为雷达图1—5分。", bold_prefix="说明：")

    add_heading(doc, "三、有效性与汇总规则", 1)
    add_bullet(doc, "单个评分人对某一维度至少完成3道有效题，该评分人的此维度结果才有效；“无法判断”不计入分母。")
    add_bullet(doc, "评分人维度得分＝该维度所有有效题百分制得分的算术平均值。")
    add_bullet(doc, "同一层级有多位评分人时，先计算每位评分人的维度得分，再计算该层级评分人的算术平均值。")
    add_bullet(doc, "默认层级权重：直属上级50%、同级30%、下属20%；没有有效评分人的层级不参与，并按现有有效层级重新归一化。")
    add_bullet(doc, "最终维度百分制得分＝各有效层级均分的加权平均值；雷达图五分值＝最终百分制得分÷20。")
    add_bullet(doc, "历史对比应使用相同被评估人、相同职位和相同评分口径；系统可按评估日期调用上一次已保存的结果。")
    add_info_box(
        doc,
        "题目编码规则",
        [
            ("格式", "职位代码-维度序号-维度内题号，例如 XS-01-01"),
            ("职位代码", "XS销售、SW商务、PM项目经理、JS技术交付、XX线下运营、IP运营、FX内容发行"),
            ("编号用途", "网页录入、题目维护、数据追踪和版本更新时，均以唯一编码为准"),
        ],
    )

    add_heading(doc, "四、岗位与六维总览", 1)
    add_overview_table(doc)
    add_body(doc, "维度结构统一为：2项通用基础能力 + 2项专业核心能力 + 2项职业素养能力。各职位使用不同维度名称，以保证评价内容贴近岗位。", bold_prefix="维度结构统一为：")

    # 七个职位章节
    for role_index, role in enumerate(ROLES, 1):
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, f"职位{role_index:02d}｜{role.name}｜{role.code}｜30题")

        add_heading(doc, f"职位{role_index:02d}｜{role.name}", 1)
        add_info_box(
            doc,
            "职位题卷信息",
            [
                ("职位代码", role.code),
                ("职位定位", role.positioning),
                ("评价周期", role.period),
                ("题目结构", "6个维度 × 每维度5题＝30题"),
            ],
        )
        add_role_fields(doc, role.name)

        for dimension_index, dimension in enumerate(role.dimensions, 1):
            add_dimension_banner(doc, dimension_index, dimension)
            for question_index, question in enumerate(dimension.questions, 1):
                add_question_block(doc, role, dimension_index, question_index, question)

    return doc


def verify_saved_document(path: Path) -> dict[str, int]:
    reopened = Document(path)
    all_text: list[str] = []
    all_text.extend(p.text for p in reopened.paragraphs)
    for table in reopened.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text.extend(p.text for p in cell.paragraphs)
    text = "\n".join(all_text)
    codes = re.findall(r"题目编码：((?:XS|SW|PM|JS|XX|IP|FX)-\d{2}-\d{2})", text)
    assert len(codes) == 210, f"Word中题目编码数量异常：{len(codes)}"
    assert len(set(codes)) == 210, "Word中存在重复题目编码"
    for role_index, role in enumerate(ROLES, 1):
        assert f"职位{role_index:02d}｜{role.name}" in text
        assert sum(1 for c in codes if c.startswith(f"{role.code}-")) == 30
        for dimension in role.dimensions:
            assert dimension.name in text
    return {
        "roles": len(ROLES),
        "dimensions": sum(len(role.dimensions) for role in ROLES),
        "questions": len(codes),
        "tables": len(reopened.tables),
        "sections": len(reopened.sections),
    }


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    result = verify_saved_document(OUTPUT)
    print(OUTPUT)
    print(result)
    print(f"size_bytes={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
